import sys
from PySide6.QtWidgets import (
    QApplication, QMessageBox, QInputDialog, QTableWidgetItem, QFileDialog, QLineEdit
)
from PySide6.QtCore import Qt, QDate
from PySide6.QtGui import QFont
import openpyxl
from openpyxl.styles import Font, PatternFill

from database import FarmaciaDB
import styles
# importamos la vista gráfica que creamos en el otro archivo
from interfaz_grafica import InterfazFarmacia 

# heredamos de InterfazFarmacia para tener todos los widgets ya creados
class FarmaciaComunalApp(InterfazFarmacia):
    def __init__(self):
        super().__init__()
        
        self.db = FarmaciaDB()
        self.setStyleSheet(styles.get_stylesheet())

        self.current_role = None
        self.current_user = None

        # conectar los botones de la interfaz gráfica a las funciones lógicas
        self.conectar_botones()

        self.actualizar_tabla()
        self.actualizar_estado_sesion()

    def conectar_botones(self):
        
        self.btn_logout.clicked.connect(self.logout)
        
        # Login y creación de usuarios
        self.btn_login_bodega.clicked.connect(lambda: self.login("bodega"))
        self.btn_nuevo_usuario.clicked.connect(self.dialogo_crear_usuario)
        self.btn_login_entrega.clicked.connect(lambda: self.login("entrega"))
        
        # ingreso bodega
        self.btn_guardar.clicked.connect(self.guardar_datos)
        
        # gestión
        self.btn_entregar.clicked.connect(self.procesar_entrega)
        self.btn_editar.clicked.connect(self.editar_medicamento)
        self.btn_eliminar.clicked.connect(self.eliminar_medicamento)
        self.btn_excel.clicked.connect(self.exportar_excel)
        self.btn_alertas.clicked.connect(self.mostrar_alertas)

    # =========================================================================
    
    def login(self, tipo_acceso):
        user, ok1 = QInputDialog.getText(self, "Inicio de Sesión", f"Usuario para {tipo_acceso}:")
        if not ok1 or not user.strip():
            return False

        password, ok2 = QInputDialog.getText(self, "Inicio de Sesión", f"Contraseña para {tipo_acceso}:", QLineEdit.Password)
        if not ok2 or not password:
            return False

        rol = self.db.validar_credenciales(user.strip(), password)

        if rol == tipo_acceso:
            self.current_role = tipo_acceso
            self.current_user = user.strip()
            self.actualizar_estado_sesion()
            QMessageBox.information(self, "Bienvenido", f"Sesión iniciada como {tipo_acceso.capitalize()}.")
            return True
        else:
            QMessageBox.critical(self, "Error al ingresar", f"Usuario o contraseña incorrectos para {tipo_acceso}.")
            self.current_role = None
            self.current_user = None
            self.actualizar_estado_sesion()
            return False

    def logout(self):
        self.current_role = None
        self.current_user = None
        self.actualizar_estado_sesion()
        QMessageBox.information(self, "Sesión Cerrada", "Ha cerrado sesión exitosamente.")

    def actualizar_estado_sesion(self):
        if self.current_role:
            self.lbl_sesion.setText(f"Sesión: {self.current_role.capitalize()} ({self.current_user})")
            self.lbl_sesion.setProperty("estado", "activa")
            self.btn_logout.setEnabled(True)
        else:
            self.lbl_sesion.setText("Sin sesión activa")
            self.lbl_sesion.setProperty("estado", "inactiva")
            self.btn_logout.setEnabled(False)

        self.lbl_sesion.style().unpolish(self.lbl_sesion)
        self.lbl_sesion.style().polish(self.lbl_sesion)

    def guardar_datos(self):
        if self.current_role != "bodega":
            QMessageBox.warning(self, "Acceso Denegado", "Debe iniciar sesión en bodega.")
            return

        nombre = self.ent_nombre.text().strip()
        cantidad = self.ent_cantidad.text().strip()
        fecha = self.ent_fecha.date().toString("yyyy-MM-dd") # Extraído del QDateEdit
        marca = self.ent_marca.text().strip()

        if not nombre or not cantidad:
            QMessageBox.warning(self, "Campos incompletos", "Llenar al menos nombre y cantidad.")
            return

        try:
            cant_int = int(cantidad)
            if cant_int <= 0:
                QMessageBox.warning(self, "Error", "La cantidad debe ser positiva.")
                return
        except ValueError:
            QMessageBox.warning(self, "Error", "La cantidad debe ser un número entero.")
            return

        self.db.registrar_medicamento(nombre, cant_int, fecha, marca)
        QMessageBox.information(self, "Éxito", "Medicamento registrado en bodega")
        self.actualizar_tabla()

        self.ent_nombre.clear()
        self.ent_cantidad.clear()
        self.ent_fecha.setDate(QDate.currentDate())
        self.ent_marca.clear()

    def actualizar_tabla(self):
        self.table.setRowCount(0)
        datos = self.db.obtener_todo()

        for row_idx, fila in enumerate(datos):
            self.table.insertRow(row_idx)
            for col_idx, valor in enumerate(fila):
                item = QTableWidgetItem(str(valor))
                item.setTextAlignment(Qt.AlignCenter)
                self.table.setItem(row_idx, col_idx, item)

    def obtener_fila_seleccionada(self):
        selected_items = self.table.selectedItems()
        if not selected_items:
            QMessageBox.warning(self, "Error", "Seleccione un medicamento.")
            return None
        row = selected_items[0].row()
        try:
            return int(self.table.item(row, 0).text())
        except (ValueError, AttributeError):
            return None

    def procesar_entrega(self):
        if self.current_role != "entrega":
            QMessageBox.warning(self, "Acceso Denegado", "Debe iniciar sesión en entrega.")
            return

        item_id = self.obtener_fila_seleccionada()
        if item_id is None: return

        medicamento = self.db.obtener_medicamento_por_id(item_id)
        if medicamento is None:
            self.actualizar_tabla()
            return

        nombre_med = medicamento[1]
        stock_actual = medicamento[2]

        cantidad, ok = QInputDialog.getInt(
            self, "Entrega",
            f"Medicamento: {nombre_med}\nStock disponible: {stock_actual}\n¿Unidades a entregar?",
            minValue=1, maxValue=max(stock_actual, 1), value=1
        )
        if not ok: return

        exito, nuevo_stock = self.db.entregar_medicamento(item_id, cantidad)
        if exito:
            QMessageBox.information(self, "Exitosa", f"Se entregaron {cantidad} de {nombre_med}.\nQuedan {nuevo_stock}.")
            self.actualizar_tabla()
        else:
            QMessageBox.critical(self, "Error", "No hay unidades suficientes.")

    def editar_medicamento(self):
        if self.current_role != "bodega":
            QMessageBox.warning(self, "Acceso Denegado", "Debe iniciar sesión en bodega.")
            return

        item_id = self.obtener_fila_seleccionada()
        if item_id is None: return
        
        medicamento = self.db.obtener_medicamento_por_id(item_id)
        if medicamento is None: return

        nombre, ok1 = QInputDialog.getText(self, "Editar", "Nombre:", text=medicamento[1])
        if not ok1 or not nombre.strip(): return

        cantidad, ok2 = QInputDialog.getInt(self, "Editar", "Cantidad:", value=medicamento[2], minValue=0)
        if not ok2: return

        fecha, ok3 = QInputDialog.getText(self, "Editar", "Fecha (YYYY-MM-DD):", text=medicamento[3])
        if not ok3: return

        marca, ok4 = QInputDialog.getText(self, "Editar", "Laboratorio:", text=medicamento[4])
        if not ok4: return

        exito = self.db.editar_medicamento(item_id, nombre.strip().upper(), cantidad, fecha, marca.strip().upper())
        if exito:
            QMessageBox.information(self, "Éxito", "Actualizado correctamente.")
            self.actualizar_tabla()

    def eliminar_medicamento(self):
        if self.current_role != "bodega":
            QMessageBox.warning(self, "Acceso Denegado", "Debe iniciar sesión en bodega.")
            return

        item_id = self.obtener_fila_seleccionada()
        if item_id is None: return

        medicamento = self.db.obtener_medicamento_por_id(item_id)
        if medicamento is None: return

        nombre_med = medicamento[1]
        respuesta = QMessageBox.question(
            self, "Confirmar", f"¿Eliminar '{nombre_med}'?\nEsta acción no se puede deshacer.",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        )

        if respuesta == QMessageBox.Yes:
            exito = self.db.eliminar_medicamento(item_id)
            if exito:
                QMessageBox.information(self, "Eliminado", f"'{nombre_med}' fue eliminado.")
                self.actualizar_tabla()

    def mostrar_alertas(self):
        try:
            alertas = self.db.alertas_criticas()
        except Exception as e:
            return

        if not alertas:
            QMessageBox.information(self, "Reporte", "No hay medicamentos por vencer o agotarse.")
            return

        mensaje = "PRODUCTOS CON POCO STOCK O POR VENCER:\n\n"
        for a in alertas:
            mensaje += f"• {a[0]} | {a[1]} unidades | Vence: {a[2]} | {a[3]}\n"

        QMessageBox.warning(self, "Alerta de Inventario", mensaje)

    def exportar_excel(self):
        datos = self.db.obtener_todo()
        if not datos: return
        
        ruta_archivo, _ = QFileDialog.getSaveFileName(self, "Guardar Reporte Excel", "Reporte_Inventario.xlsx", "Archivos Excel (*.xlsx)")
        if not ruta_archivo: return

        try:
            libro = openpyxl.Workbook()
            hoja = libro.active
            hoja.title = "Inventario"
            encabezados = ["ID", "Medicamento", "Unidades Disponibles", "Fecha de Vencimiento", "Laboratorio"]
            hoja.append(encabezados)

            relleno = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            fuente = Font(bold=True, color="FFFFFF")
            
            for celda in hoja[1]:
                celda.font = fuente
                celda.fill = relleno

            for fila in datos: hoja.append(fila)

            hoja.column_dimensions['B'].width = 30
            hoja.column_dimensions['C'].width = 25
            hoja.column_dimensions['D'].width = 25
            hoja.column_dimensions['E'].width = 25

            libro.save(ruta_archivo)
            QMessageBox.information(self, "Éxito", "Reporte Excel generado correctamente.")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo generar el Excel:\n{e}")

    def dialogo_crear_usuario(self):
        roles = ["bodega", "entrega"]
        rol, ok1 = QInputDialog.getItem(self, "Nuevo Usuario", "Seleccione el rol:", roles, 0, False)
        if not ok1 or not rol: return

        usuario, ok2 = QInputDialog.getText(self, "Nuevo Usuario", f"Nombre de usuario para {rol}:")
        if not ok2 or not usuario.strip(): return

        password, ok3 = QInputDialog.getText(self, "Nuevo Usuario", "Contraseña:", QLineEdit.Password)
        if not ok3 or not password: return

        exito, mensaje = self.db.agregar_usuario(rol, usuario.strip(), password)
        if exito:
            QMessageBox.information(self, "Éxito", mensaje)
        else:
            QMessageBox.critical(self, "Error", mensaje)

    def closeEvent(self, event):
        self.db.close()
        event.accept()

# entrada  del programa
if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setFont(QFont("Segoe UI", 10))
    ventana = FarmaciaComunalApp()
    ventana.show()
    sys.exit(app.exec())